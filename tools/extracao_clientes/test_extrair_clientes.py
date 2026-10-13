import csv
import tempfile
import unittest
from pathlib import Path
from argparse import Namespace

from docx import Document
from extrair_clientes import parse_qualification, extract_candidates, consolidate, run, FIELDS


def qualification(name='Maria Exemplo Silva', cpf='001.234.567-89', gender='brasileira, solteira', rg='01.234.567-X', address='Rua 25 de Março, Nº 75 casa 2, Vila Exemplo, Cep 13338-190, São José — SP'):
    return f'{name}, nacionalidade: {gender}, profissão: atendente, carteira de identidade (RG) Nº {rg}, expedida por SSP/SP, CPF Nº {cpf}, endereço: {address}'


class ParserTests(unittest.TestCase):
    def test_feminine_and_address(self):
        r = parse_qualification(qualification())
        expected = dict(nomeCompleto='Maria Exemplo Silva', nacionalidade='brasileira', estadoCivil='SOLTEIRO', sexoDocumento='FEMININO', profissao='atendente', rg='01.234.567-X', rgOrgaoExpedidor='SSP/SP', cpf='001.234.567-89', enderecoRua='Rua 25 de Março', enderecoNumero='75', enderecoComplemento='casa 2', enderecoBairro='Vila Exemplo', enderecoCep='13338-190', enderecoCidade='São José', enderecoEstado='SP', statusExtracao='OK')
        for field, value in expected.items():
            self.assertEqual(r[field], value, field)

    def test_variants(self):
        for number in ('Nº', 'nº', 'N°', 'n°', 'No', 'Nro'):
            with self.subTest(number=number):
                r = parse_qualification(qualification(gender='brasileiro, casado', rg='01234567X', cpf='00123456789', address=f'Rua 7 de Setembro, {number} 9, CEP 01234567, Rio Claro - SP').replace('Nº', number).replace('expedida', 'expedido'))
                self.assertEqual(r['sexoDocumento'], 'MASCULINO')
                self.assertEqual(r['estadoCivil'], 'CASADO')
                self.assertEqual(r['rg'], '01.234.567-X')
                self.assertEqual(r['enderecoCep'], '01234-567')
                self.assertEqual(r['enderecoNumero'], '9')
                self.assertEqual(r['enderecoBairro'], '')
                self.assertEqual(r['enderecoComplemento'], '')

    def test_civil_and_gender_ambiguity(self):
        for civil, expected in [('união estável', 'UNIAO_ESTAVEL'), ('divorciada', 'DIVORCIADO'), ('viúva', 'VIUVO'), ('separada', 'SEPARADO')]:
            self.assertEqual(parse_qualification(qualification(gender='brasileira, '+civil))['estadoCivil'], expected)
        r = parse_qualification(qualification(gender='brasileira, casado'))
        self.assertEqual(r['sexoDocumento'], '')
        self.assertEqual(r['statusExtracao'], 'REVISAR')
        r = parse_qualification(qualification(gender='portuguesa, união estável'))
        self.assertEqual(r['sexoDocumento'], '')

    def test_creditor_and_guarantor(self):
        ps = [qualification(), 'Doravante denominado “DEVEDOR”,', qualification(name='Edson Antonio Vieira', cpf='072.848.058-11'), 'Doravante denominado “CREDOR”', qualification(name='Joana Exemplo Santos', cpf='123.456.789-01'), 'Doravante denominada “FIADORA”']
        rows, error = extract_candidates(ps)
        self.assertFalse(error)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['nomeCompleto'], 'Maria Exemplo Silva')
        self.assertIn('Fiador encontrado', rows[0]['observacoes'])

    def test_ambiguous_people(self):
        rows, error = extract_candidates([qualification(), qualification(name='Outra Pessoa Exemplo', cpf='12345678901')])
        self.assertEqual(rows, [])
        self.assertTrue(error)

    def test_consolidation_and_global_duplicate(self):
        older = parse_qualification(qualification())
        older.update(pastaOrigem='A', arquivoOrigem='old.docx', dataArquivo='2023', _mtime=1, _docx=True)
        newer = parse_qualification(qualification(address='Rua Nova, Nº 4, CEP 13338190, Rio Claro - SP'))
        newer.update(pastaOrigem='B', arquivoOrigem='new.docx', dataArquivo='2026', _mtime=2, _docx=True)
        rows, sources = consolidate([older, newer])
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['enderecoRua'], 'Rua Nova')
        self.assertEqual(rows[0]['enderecoComplemento'], '')
        self.assertEqual(rows[0]['statusExtracao'], 'DUPLICIDADE')
        self.assertIn('new.docx', sources)

    def test_without_cpf_not_dedup_by_name(self):
        rows = []
        for i, rg in enumerate(('1234567', '7654321')):
            r = parse_qualification(qualification(cpf='', rg=rg))
            r.update(pastaOrigem='A', arquivoOrigem=str(i), _mtime=i, _docx=True)
            rows.append(r)
        self.assertEqual(len(consolidate(rows)[0]), 2)

    def test_end_to_end_partial_failure_and_snapshots(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            folder = root/'Maria Exemplo Silva'
            folder.mkdir()
            nested = folder/'historico'
            nested.mkdir()
            doc = Document()
            doc.add_paragraph(qualification())
            doc.add_paragraph('Doravante denominado “DEVEDOR”')
            path = nested/'contrato.docx'
            doc.save(path)
            original = path.read_bytes()
            (folder/'corrompido.docx').write_bytes(b'broken')
            (root/'Sem documentos').mkdir()
            args = Namespace(raiz=root, saida=None, limite=None, pasta=None, paragrafos=20, libreoffice=None)
            summary = run(args)
            self.assertEqual(summary['clientes_unicos'], 1)
            self.assertEqual(summary['erros_leitura'], 1)
            self.assertEqual(summary['pastas_sem_word'], 1)
            self.assertEqual(path.read_bytes(), original)
            out = root/'_export_clientes'
            self.assertTrue((out/'clientes_extraidos.csv').read_bytes().startswith(b'\xef\xbb\xbf'))
            with (out/'clientes_extraidos.csv').open(encoding='utf-8-sig', newline='') as f:
                rows = list(csv.DictReader(f, delimiter=';'))
            self.assertEqual(rows[0]['cpf'], '001.234.567-89')
            self.assertEqual(list(rows[0]), FIELDS)
            run(args)
            self.assertEqual(len(list(out.glob('clientes_extraidos_*.csv'))), 2)


if __name__ == '__main__':
    unittest.main()
